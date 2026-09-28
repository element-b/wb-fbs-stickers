from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def make_mysklad_xlsx(groups: list[dict]) -> bytes:
    """
    Создаёт XLSX-файл для работы с МойСклад.

    Столбцы:

    - Артикул:
      один раз для каждого уникального артикула;

    - Количество:
      количество сборочных заданий / оригинальных стикеров WB
      для соответствующего артикула;

    - СЦ назначения:
      независимый список уникальных СЦ из выбранных поставок.
      Список нужен для ручного копирования в комментарий отгрузки.

    Третий столбец не является построчной связью с артикулом.
    """
    if not groups:
        raise ValueError(
            "Нельзя создать XLSX: отсутствуют группы артикулов."
        )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "WB FBS"

    header_fill = PatternFill(
        start_color="009B77",
        end_color="009B77",
        fill_type="solid",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    worksheet.append(
        [
            "Артикул",
            "Количество",
            "СЦ назначения",
        ]
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    article_rows: list[tuple[str, int]] = []
    unique_centers: list[str] = []
    seen_centers: set[str] = set()

    for group in groups:
        article = str(group.get("article") or "").strip()
        order_ids = group.get("order_ids", [])
        centers = group.get("distribution_centers", [])

        if not article:
            raise ValueError(
                "Нельзя создать XLSX: найден пустой артикул."
            )

        if not isinstance(order_ids, list) or not order_ids:
            raise ValueError(
                f"Нельзя создать XLSX: отсутствуют задания "
                f"у артикула `{article}`."
            )

        # Одна строка на артикул, количество привязано к нему.
        article_rows.append(
            (
                article,
                len(order_ids),
            )
        )

        if not isinstance(centers, list):
            centers = []

        for center in centers:
            center_text = str(center).strip()

            if not center_text:
                continue

            # СЦ уникализируются без учёта регистра,
            # но в XLSX сохраняется исходное написание.
            center_key = center_text.casefold()

            if center_key in seen_centers:
                continue

            seen_centers.add(center_key)
            unique_centers.append(center_text)

    if not article_rows:
        raise ValueError(
            "Нельзя создать XLSX без строк с артикулами."
        )

    # Первые два столбца — список артикулов с количеством.
    # Третий столбец — независимый вертикальный список уникальных СЦ.
    max_rows = max(
        len(article_rows),
        len(unique_centers),
    )

    for index in range(max_rows):
        if index < len(article_rows):
            article, quantity = article_rows[index]
        else:
            article = ""
            quantity = ""

        center = (
            unique_centers[index]
            if index < len(unique_centers)
            else ""
        )

        worksheet.append(
            [
                article,
                quantity,
                center,
            ]
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    worksheet.column_dimensions["A"].width = 42
    worksheet.column_dimensions["B"].width = 16
    worksheet.column_dimensions["C"].width = 34

    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=3,
    ):
        row[0].alignment = Alignment(
            horizontal="left",
            vertical="center",
        )

        row[1].alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        row[2].alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True,
        )

    output = BytesIO()
    workbook.save(output)

    return output.getvalue()