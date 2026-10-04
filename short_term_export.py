from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


HEADER_FILL = PatternFill(
    start_color="009B77",
    end_color="009B77",
    fill_type="solid",
)

HEADER_FONT = Font(
    color="FFFFFF",
    bold=True,
)

CRITICAL_FILL = PatternFill(
    start_color="FEE2E2",
    end_color="FEE2E2",
    fill_type="solid",
)

URGENT_FILL = PatternFill(
    start_color="FFEDD5",
    end_color="FFEDD5",
    fill_type="solid",
)

REPLENISH_FILL = PatternFill(
    start_color="FEF9C3",
    end_color="FEF9C3",
    fill_type="solid",
)

NORMAL_FILL = PatternFill(
    start_color="DCFCE7",
    end_color="DCFCE7",
    fill_type="solid",
)


def _priority_fill(value: object) -> PatternFill | None:
    """Возвращает цвет строки XLSX по приоритету."""
    text = str(value or "")

    if text.startswith("🔴"):
        return CRITICAL_FILL

    if text.startswith("🟠"):
        return URGENT_FILL

    if text.startswith("🟡"):
        return REPLENISH_FILL

    if text.startswith("🟢"):
        return NORMAL_FILL

    return None


def _append_dataframe(
    worksheet,
    dataframe: pd.DataFrame,
) -> None:
    """Записывает таблицу на лист XLSX и оформляет её."""
    columns = dataframe.columns.tolist()

    worksheet.append(columns)

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for row in dataframe.to_dict(orient="records"):
        worksheet.append(
            [
                row.get(column, "")
                for column in columns
            ]
        )

        fill = _priority_fill(row.get("Приоритет"))

        if fill is not None:
            for cell in worksheet[worksheet.max_row]:
                cell.fill = fill

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False

    for column_cells in worksheet.columns:
        max_length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )

        width = min(
            max(max_length + 2, 12),
            44,
        )

        worksheet.column_dimensions[
            column_cells[0].column_letter
        ].width = width

    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row,
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )


def make_short_term_plan_xlsx(
    queue_table: pd.DataFrame,
    full_table: pd.DataFrame,
    parameters: dict[str, str | int],
) -> bytes:
    """
    Создаёт XLSX краткосрочного плана WB FBS.

    Листы:

    - Очередь производства;
    - Все артикулы;
    - Параметры.
    """
    if full_table.empty:
        raise ValueError(
            "Нельзя создать XLSX: отсутствуют строки плана."
        )

    workbook = Workbook()

    queue_sheet = workbook.active
    queue_sheet.title = "Очередь производства"

    _append_dataframe(
        worksheet=queue_sheet,
        dataframe=queue_table,
    )

    full_sheet = workbook.create_sheet(
        title="Все артикулы"
    )

    _append_dataframe(
        worksheet=full_sheet,
        dataframe=full_table,
    )

    parameters_sheet = workbook.create_sheet(
        title="Параметры"
    )

    parameters_sheet.append(
        [
            "Параметр",
            "Значение",
        ]
    )

    for cell in parameters_sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for key, value in parameters.items():
        parameters_sheet.append(
            [
                key,
                value,
            ]
        )

    parameters_sheet.freeze_panes = "A2"
    parameters_sheet.sheet_view.showGridLines = False
    parameters_sheet.column_dimensions["A"].width = 40
    parameters_sheet.column_dimensions["B"].width = 55

    for row in parameters_sheet.iter_rows(
        min_row=2,
        max_row=parameters_sheet.max_row,
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

    output = BytesIO()
    workbook.save(output)

    return output.getvalue()