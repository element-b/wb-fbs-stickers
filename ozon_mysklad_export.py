from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


HEADER_FILL = PatternFill(
    start_color="009B77",
    end_color="009B77",
    fill_type="solid",
)

HEADER_FONT = Font(
    color="FFFFFF",
    bold=True,
)


def make_ozon_mysklad_xlsx(
    article_rows: list[dict],
) -> bytes:
    """
    Создаёт XLSX для импорта или ручной работы в МоемСкладе.

    В файле только два связанных столбца:

    - Артикул;
    - Количество.

    Количество — фактическая сумма `quantity` товаров с offer_id
    соответствующего артикула `NAKL_` в выбранных FBS-отправлениях
    Ozon.

    Номера отправлений намеренно не добавляются в основной лист:
    они не требуются для импорта товарных строк в МойСклад.
    """
    if not article_rows:
        raise ValueError(
            "Нельзя создать XLSX: отсутствуют артикулы Ozon."
        )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Ozon FBS"

    worksheet.append(
        [
            "Артикул",
            "Количество",
        ]
    )

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    seen_articles: set[str] = set()

    for row in article_rows:
        if not isinstance(row, dict):
            raise ValueError(
                "Нельзя создать XLSX: найдена некорректная строка."
            )

        article = str(
            row.get("Артикул Ozon") or ""
        ).strip()

        if not article:
            raise ValueError(
                "Нельзя создать XLSX: найден пустой артикул."
            )

        article_key = article.casefold()

        if article_key in seen_articles:
            raise ValueError(
                "Нельзя создать XLSX: один артикул Ozon повторяется "
                "в итоговой сводке."
            )

        seen_articles.add(article_key)

        quantity = row.get("Количество товаров")

        if isinstance(quantity, bool):
            raise ValueError(
                f"Некорректное количество у артикула `{article}`."
            )

        try:
            normalized_quantity = int(quantity)
        except (TypeError, ValueError) as error:
            raise ValueError(
                f"Некорректное количество у артикула `{article}`."
            ) from error

        if normalized_quantity <= 0:
            raise ValueError(
                f"Количество у артикула `{article}` должно быть "
                "положительным."
            )

        worksheet.append(
            [
                article,
                normalized_quantity,
            ]
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False

    worksheet.column_dimensions["A"].width = 42
    worksheet.column_dimensions["B"].width = 16

    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=2,
    ):
        row[0].alignment = Alignment(
            horizontal="left",
            vertical="center",
        )

        row[1].alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    output = BytesIO()
    workbook.save(output)

    return output.getvalue()